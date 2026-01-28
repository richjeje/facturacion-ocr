import cv2
import pytesseract
import numpy as np
from pdf2image import convert_from_path
import pdfplumber
from PIL import Image
import os
import logging
import time
import easyocr
from docx import Document
from pptx import Presentation
from core.utils import ExtractionError
from concurrent.futures import ThreadPoolExecutor

logger = logging.getLogger(__name__)

# Global reader instance for EasyOCR (Lazy loading)
_easyocr_reader = None


def get_easyocr_reader():
    global _easyocr_reader
    if _easyocr_reader is None:
        logger.info("Inicializando EasyOCR Reader (GPU si está disponible)...")
        _easyocr_reader = easyocr.Reader(["es", "en"])
    return _easyocr_reader


def extract_text_from_docx(file_path):
    """Extrae texto de un archivo DOCX."""
    try:
        doc = Document(file_path)
        texto = ""
        for para in doc.paragraphs:
            texto += para.text + "\n"
        logger.info(f"Texto extraído de DOCX: {file_path}")
        return texto, "DOCX"
    except Exception as e:
        logger.error(f"Error extrayendo texto de DOCX {file_path}: {e}")
        raise ExtractionError(f"Error en DOCX {file_path}: {e}")


def extract_text_from_pptx(file_path):
    """Extrae texto de un archivo PPTX."""
    try:
        prs = Presentation(file_path)
        texto = ""
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    texto += shape.text + "\n"
        logger.info(f"Texto extraído de PPTX: {file_path}")
        return texto, "PPTX"
    except Exception as e:
        logger.error(f"Error extrayendo texto de PPTX {file_path}: {e}")
        raise ExtractionError(f"Error en PPTX {file_path}: {e}")


def extract_text_from_pdf(file_path):
    """Extrae texto de un PDF usando pdfplumber, con fallback a OCR si no hay texto."""
    try:
        with pdfplumber.open(file_path) as pdf:
            texto = ""
            for page in pdf.pages:
                page_text = page.extract_text() or ""
                texto += page_text + "\n"
        if texto.strip():
            logger.info(f"Texto extraído exitosamente de PDF: {file_path}")
            return texto, "Texto"
        else:
            logger.info(f"No texto en PDF {file_path}, intentando OCR...")
            return _ocr_from_pdf(file_path)
    except Exception as e:
        logger.error(f"Error extrayendo texto de PDF {file_path}: {e}")
        raise ExtractionError(f"Error en PDF {file_path}: {e}")


def _ocr_from_pdf(file_path):
    """Realiza OCR en PDF con reintentos y procesamiento paralelo de páginas."""
    max_retries = 2
    for attempt in range(max_retries + 1):
        try:
            images = convert_from_path(file_path)
            
            def process_page(img):
                img_cv = cv2.cvtColor(np.array(img), cv2.COLOR_RGB2GRAY)
                img_cv = cv2.threshold(
                    img_cv, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU
                )[1]
                return pytesseract.image_to_string(img_cv, lang="spa+eng")

            with ThreadPoolExecutor(max_workers=min(len(images), 4)) as executor:
                page_texts = list(executor.map(process_page, images))
            
            texto = "\n".join(page_texts)
            if texto.strip():
                logger.info(f"OCR exitoso en PDF {file_path} en intento {attempt + 1}")
                return texto, "OCR"
        except Exception as e:
            logger.warning(f"Error en OCR de PDF {file_path}, intento {attempt + 1}: {e}")
            if attempt < max_retries:
                time.sleep(1 * (attempt + 1))
    raise ExtractionError(f"OCR falló en PDF {file_path}")


def extract_text_from_image(file_path):
    """Extrae texto de una imagen usando EasyOCR con reintentos y optimización."""
    max_retries = 2
    for attempt in range(max_retries + 1):
        try:
            img_pil = Image.open(file_path)
            if img_pil.width > 2000 or img_pil.height > 2000:
                ratio = min(2000 / img_pil.width, 2000 / img_pil.height)
                new_size = (int(img_pil.width * ratio), int(img_pil.height * ratio))
                img_pil = img_pil.resize(new_size, Image.Resampling.LANCZOS)

            img_array = np.array(img_pil)
            reader = get_easyocr_reader()
            results = reader.readtext(img_array, detail=0, paragraph=True)
            texto = " ".join(results)
            if texto.strip():
                return texto, "EasyOCR"
        except Exception as e:
            if attempt < max_retries:
                time.sleep(1 * (attempt + 1))
    raise ExtractionError(f"EasyOCR falló en imagen {file_path}")


def extract_text(file_path):
    """Función principal para extraer texto según tipo de archivo."""
    if not os.path.exists(file_path):
        raise ExtractionError(f"Archivo no existe: {file_path}")
    file_size = os.path.getsize(file_path)
    if file_size > 50 * 1024 * 1024:
        raise ExtractionError(f"Archivo demasiado grande: {file_path}")
    
    ext = file_path.lower().split(".")[-1]
    if ext == "pdf":
        return extract_text_from_pdf(file_path)
    elif ext in ["jpg", "jpeg", "png"]:
        return extract_text_from_image(file_path)
    elif ext == "docx":
        return extract_text_from_docx(file_path)
    elif ext == "pptx":
        return extract_text_from_pptx(file_path)
    else:
        raise ExtractionError(f"Tipo de archivo no soportado: {ext}")
